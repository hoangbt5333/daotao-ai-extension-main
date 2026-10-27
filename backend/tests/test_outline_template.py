"""
Tests for Course Outline Excel Template generation and error handling.
"""
import sys
from io import BytesIO
from pathlib import Path
import pytest
from openpyxl import load_workbook
from fastapi import UploadFile

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from app.api.v1.endpoints.outline import (
    REQUIRED_HEADERS,
    OPTIONAL_HEADERS,
    TEMPLATE_HEADERS,
    SUPPORTED_UNIT_TYPES,
    download_outline_template,
    import_outline_excel,
)
from app.config.settings import get_settings


# ===========================================================================
# Task 1 Tests: Template Structure and Clean Import
# ===========================================================================

@pytest.mark.anyio
async def test_download_outline_template_structure():
    """Endpoint download_outline_template phải sinh file .xlsx với đúng cấu trúc:
    - 2 sheets: 'Mẫu' và 'Hướng dẫn'.
    - Sheet 'Mẫu': Header khớp TEMPLATE_HEADERS, đủ 1 dòng cho mỗi SUPPORTED_UNIT_TYPES,
      và 1 dòng để trống Type nhưng điền Slide để kiểm tra tự nhận diện.
    - Sheet 'Hướng dẫn': Mô tả các cột, danh sách Type hợp lệ, quy tắc đặt tên,
      giới hạn kích thước lấy từ settings.file.max_file_size (50MB).
    """
    resp = await download_outline_template()
    assert resp.status_code == 200
    assert resp.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="outline_template.xlsx"' in resp.headers.get("Content-Disposition", "")

    # Đọc lại bằng openpyxl
    wb = load_workbook(BytesIO(resp.body))
    assert wb.sheetnames == ["Mẫu", "Hướng dẫn"]

    # 1. Kiểm tra sheet 'Mẫu'
    ws_mau = wb["Mẫu"]
    headers = [cell.value for cell in ws_mau[1] if cell.value is not None]
    assert headers == TEMPLATE_HEADERS
    assert headers == ["Chapter", "Lesson", "Unit", "Type", "Content", "Slide", "Youtube"]

    rows = list(ws_mau.iter_rows(min_row=2, values_only=True))
    # Phải có ít nhất len(SUPPORTED_UNIT_TYPES) + 1 dòng (cho auto-detect)
    assert len(rows) >= len(SUPPORTED_UNIT_TYPES) + 1

    # Kiểm tra sự xuất hiện của các types được hỗ trợ
    types_in_sample = [r[headers.index("Type")] for r in rows if r[headers.index("Type")]]
    for expected_type in SUPPORTED_UNIT_TYPES:
        assert expected_type in types_in_sample, f"Thiếu ví dụ cho Type '{expected_type}' trong sheet Mẫu"

    # Kiểm tra có 1 dòng không điền Type nhưng có điền Slide hoặc Youtube (auto-detection)
    auto_detect_rows = [r for r in rows if not r[headers.index("Type")] and (r[headers.index("Slide")] or r[headers.index("Youtube")])]
    assert len(auto_detect_rows) >= 1, "Cần có ít nhất 1 dòng để trống Type nhưng có Slide/Youtube để minh họa tự nhận diện"

    # 2. Kiểm tra sheet 'Hướng dẫn'
    ws_guide = wb["Hướng dẫn"]
    guide_text = "\n".join(" ".join(str(c) for c in row if c is not None) for row in ws_guide.iter_rows(values_only=True))

    # Cột bắt buộc & tùy chọn
    for h in REQUIRED_HEADERS:
        assert h in guide_text
    for h in OPTIONAL_HEADERS:
        assert h in guide_text

    # Danh sách Type hợp lệ
    for expected_type in SUPPORTED_UNIT_TYPES:
        assert expected_type in guide_text

    # Quy tắc đặt tên file (docs 3.4.2: không dấu, không khoảng trắng)
    assert "dấu" in guide_text.lower() or "khoảng trắng" in guide_text.lower() or "ký tự" in guide_text.lower()

    # Giới hạn dung lượng từ cấu hình settings
    settings = get_settings()
    max_mb = settings.file.max_file_size // (1024 * 1024)
    assert f"{max_mb}MB" in guide_text or f"{max_mb} MB" in guide_text or str(settings.file.max_file_size) in guide_text

    # Ghi chú tương thích ngược 5 cột
    assert "5 cột" in guide_text or "tương thích" in guide_text.lower()


@pytest.mark.anyio
async def test_template_file_imports_cleanly():
    """Đưa chính file template sinh ra vào import_outline_excel():
    - Phải thành công (success=True).
    - warnings phải rỗng (không thiếu Lesson, không có cảnh báo).
    - Các unit nhận đúng type mong đợi.
    """
    resp = await download_outline_template()
    upload = UploadFile(filename="outline_template.xlsx", file=BytesIO(resp.body))

    res = await import_outline_excel(upload)
    assert res.success is True, f"Import thất bại: {res.message}"
    assert len(res.warnings) == 0, f"Có cảnh báo không mong muốn: {res.warnings}"

    outline = res.outline.get("outline", [])
    assert len(outline) > 0

    # Gom tất cả các unit từ mọi chapter và subsection
    all_units = []
    for chapter in outline:
        for sub in chapter.get("subsections", []):
            for unit in sub.get("units", []):
                all_units.append(unit)

    # Đảm bảo mỗi unit tương ứng với từng SUPPORTED_UNIT_TYPES đều được tạo ra đúng type
    unit_types = {u.get("type") for u in all_units}
    for expected_type in SUPPORTED_UNIT_TYPES:
        assert expected_type in unit_types, f"Unit type '{expected_type}' không được tạo đúng sau khi import"


# ===========================================================================
# Task 2 Tests: Backward Compatibility & Error Message Hints
# ===========================================================================

def _create_xlsx_bytes(headers: list, rows: list) -> BytesIO:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Outline"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@pytest.mark.anyio
async def test_legacy_5_columns_compatibility_xlsx():
    """File Excel cũ chỉ có 5 cột (Chapter, Lesson, Unit, Slide, Youtube)
    vẫn hoạt động đúng, tự động phân loại video/pdf/video_pdf."""
    headers = ["Chapter", "Lesson", "Unit", "Slide", "Youtube"]
    rows = [
        ["Chương 1", "Bài 1", "Unit Video", "", "https://youtu.be/abc"],
        ["Chương 1", "Bài 1", "Unit Slide", "bai1.pdf", ""],
        ["Chương 1", "Bài 1", "Unit VideoSlide", "bai1.pdf", "https://youtu.be/abc"],
        ["Chương 1", "Bài 1", "Unit Empty", "", ""],
    ]
    buf = _create_xlsx_bytes(headers, rows)
    upload = UploadFile(filename="legacy_5_columns.xlsx", file=buf)
    res = await import_outline_excel(upload)

    assert res.success is True
    units = res.outline["outline"][0]["subsections"][0]["units"]
    assert len(units) == 4
    assert units[0]["type"] == "video"
    assert units[1]["type"] == "pdf"
    assert units[2]["type"] == "video_pdf"


@pytest.mark.anyio
async def test_legacy_5_columns_compatibility_csv():
    """File CSV cũ 5 cột vẫn hoạt động tương thích ngược 100%."""
    csv_content = (
        "Chapter,Lesson,Unit,Slide,Youtube\n"
        "Chương 1,Bài 1,Unit Video,,https://youtu.be/abc\n"
        "Chương 1,Bài 1,Unit Slide,bai1.pdf,\n"
    ).encode("utf-8")
    upload = UploadFile(filename="legacy.csv", file=BytesIO(csv_content))
    res = await import_outline_excel(upload)

    assert res.success is True
    units = res.outline["outline"][0]["subsections"][0]["units"]
    assert units[0]["type"] == "video"
    assert units[1]["type"] == "pdf"


@pytest.mark.anyio
async def test_missing_required_headers_contains_hint_xlsx():
    """Khi thiếu cột bắt buộc trong .xlsx, thông báo lỗi phải nêu rõ cột thiếu
    và kèm gợi ý xem file mẫu."""
    # Thiếu cột Unit
    headers = ["Chapter", "Lesson", "Slide", "Youtube"]
    rows = [["Chương 1", "Bài 1", "bai1.pdf", ""]]
    buf = _create_xlsx_bytes(headers, rows)
    upload = UploadFile(filename="missing_header.xlsx", file=buf)
    res = await import_outline_excel(upload)

    assert res.success is False
    assert "Thiếu cột bắt buộc" in res.message
    assert "Unit" in res.message
    assert "file mẫu" in res.message.lower()


@pytest.mark.anyio
async def test_missing_required_headers_contains_hint_csv():
    """Khi thiếu cột bắt buộc trong .csv, thông báo lỗi phải nêu rõ cột thiếu
    và kèm gợi ý xem file mẫu."""
    csv_content = "Chapter,Slide,Youtube\nChương 1,bai1.pdf,\n".encode("utf-8")
    upload = UploadFile(filename="missing_header.csv", file=BytesIO(csv_content))
    res = await import_outline_excel(upload)

    assert res.success is False
    assert "Thiếu cột bắt buộc" in res.message
    assert "file mẫu" in res.message.lower()


@pytest.mark.anyio
async def test_invalid_type_error_message_contains_hint_xlsx():
    """Khi giá trị cột Type trong file .xlsx không hợp lệ, thông báo lỗi phải chỉ rõ
    hàng vi phạm, giá trị sai, gợi ý xem file mẫu và liệt kê các Type hợp lệ."""
    headers = ["Chapter", "Lesson", "Unit", "Type", "Content"]
    rows = [
        ["Chương 1", "Bài 1", "Unit 1", "invalid_unit_type", "some_content"]
    ]
    buf = _create_xlsx_bytes(headers, rows)
    upload = UploadFile(filename="invalid_type.xlsx", file=buf)
    res = await import_outline_excel(upload)

    assert res.success is False
    assert "invalid_unit_type" in res.message
    assert "file mẫu" in res.message.lower()
    for supported in SUPPORTED_UNIT_TYPES:
        assert supported in res.message, f"Thông báo lỗi cần liệt kê type '{supported}'"


@pytest.mark.anyio
async def test_invalid_type_error_message_contains_hint_csv():
    """Khi giá trị cột Type trong file .csv không hợp lệ, thông báo lỗi phải chỉ rõ
    hàng vi phạm, giá trị sai, gợi ý xem file mẫu và liệt kê các Type hợp lệ."""
    csv_content = (
        "Chapter,Lesson,Unit,Type,Content\n"
        "Chương 1,Bài 1,Unit 1,bad_type_xyz,some_content\n"
    ).encode("utf-8")
    upload = UploadFile(filename="invalid_type.csv", file=BytesIO(csv_content))
    res = await import_outline_excel(upload)

    assert res.success is False
    assert "bad_type_xyz" in res.message
    assert "file mẫu" in res.message.lower()
    for supported in SUPPORTED_UNIT_TYPES:
        assert supported in res.message, f"Thông báo lỗi cần liệt kê type '{supported}'"


# ===========================================================================
# Forward-Fill Lesson (Kế thừa Lesson trống trong cùng Chapter)
# ===========================================================================

@pytest.mark.anyio
async def test_forward_fill_lesson_and_chapter_xlsx():
    """Khi để trống ô Chapter và/hoặc Lesson ở các dòng tiếp theo (như ảnh người dùng cung cấp),
    hệ thống phải tự động kế thừa Lesson gần nhất trong cùng Chapter,
    KHÔNG sinh subsection 'General' và KHÔNG sinh warning."""
    headers = ["Chapter", "Lesson", "Unit", "Type", "Content"]
    rows = [
        ["Chương 1: Mở đầu", "Bài 1.1: Giới thiệu chung", "1.1.1 Video giới thiệu tổng quan", "video", "https://youtu.be/a"],
        ["", "", "1.1.2 Tài liệu hướng dẫn học tập", "pdf", "tailieu.pdf"],
        ["", "Bài 1.2: Luyện nghe và phát âm", "1.2.1 Bài nghe hội thoại mẫu", "audio", "nghe.mp3"],
        ["", "", "1.2.2 Tóm tắt", "text", "<p>Tom tat</p>"],
        ["Chương 2: Kiến thức trọng tâm", "Bài 2.1: Lý thuyết căn bản", "2.1.1 Bài giảng lý thuyết kèm slide", "video_pdf", "https://youtu.be/b"],
        ["", "Bài 2.2: Tài liệu tương tác trực tuyến", "2.2.1 Tài liệu nhúng Google Drive", "embed", "https://drive.google.com/test"],
        ["Chương 3: Tổng kết & Ôn tập", "Bài 3.1: Hướng dẫn ôn tập", "3.1.1 Tóm tắt kiến thức cốt lõi", "text", "<p>Cot loi</p>"],
        ["", "Bài 3.1: Hướng dẫn ôn tập", "3.1.2 Tài liệu tham khảo bổ sung", "pdf", "thamkhao.pdf"],
    ]
    buf = _create_xlsx_bytes(headers, rows)
    upload = UploadFile(filename="fill_down_sample.xlsx", file=buf)
    res = await import_outline_excel(upload)

    assert res.success is True
    assert len(res.warnings) == 0, f"Không được có cảnh báo thiếu Lesson: {res.warnings}"

    chapters = res.outline["outline"]
    assert len(chapters) == 3

    # Chương 1 phải có đúng 2 bài (1.1 và 1.2), không có bài 'General'
    c1_subs = chapters[0]["subsections"]
    assert len(c1_subs) == 2
    assert c1_subs[0]["title"] == "Bài 1.1: Giới thiệu chung"
    assert len(c1_subs[0]["units"]) == 2
    assert c1_subs[0]["units"][0]["title"] == "1.1.1 Video giới thiệu tổng quan"
    assert c1_subs[0]["units"][1]["title"] == "1.1.2 Tài liệu hướng dẫn học tập"

    assert c1_subs[1]["title"] == "Bài 1.2: Luyện nghe và phát âm"
    assert len(c1_subs[1]["units"]) == 2
    assert c1_subs[1]["units"][0]["title"] == "1.2.1 Bài nghe hội thoại mẫu"
    assert c1_subs[1]["units"][1]["title"] == "1.2.2 Tóm tắt"

    # Chương 2 phải có 2 bài (2.1 và 2.2)
    c2_subs = chapters[1]["subsections"]
    assert len(c2_subs) == 2
    assert c2_subs[0]["title"] == "Bài 2.1: Lý thuyết căn bản"
    assert len(c2_subs[0]["units"]) == 1
    assert c2_subs[1]["title"] == "Bài 2.2: Tài liệu tương tác trực tuyến"
    assert len(c2_subs[1]["units"]) == 1

    # Chương 3 phải có 1 bài (3.1) chứa cả 2 unit
    c3_subs = chapters[2]["subsections"]
    assert len(c3_subs) == 1
    assert c3_subs[0]["title"] == "Bài 3.1: Hướng dẫn ôn tập"
    assert len(c3_subs[0]["units"]) == 2


@pytest.mark.anyio
async def test_forward_fill_lesson_and_chapter_csv():
    """Kiểm tra tính năng kế thừa Lesson trống cho cả file CSV."""
    csv_content = (
        "Chapter,Lesson,Unit,Type,Content\n"
        "Chương 1,Bài 1,Unit 1.1,video,https://youtu.be/a\n"
        ",,Unit 1.2,pdf,slide.pdf\n"
        ",Bài 2,Unit 2.1,text,<p>A</p>\n"
        ",,Unit 2.2,text,<p>B</p>\n"
    ).encode("utf-8")
    upload = UploadFile(filename="fill_down.csv", file=BytesIO(csv_content))
    res = await import_outline_excel(upload)

    assert res.success is True
    assert len(res.warnings) == 0
    subs = res.outline["outline"][0]["subsections"]
    assert len(subs) == 2
    assert subs[0]["title"] == "Bài 1"
    assert len(subs[0]["units"]) == 2
    assert subs[1]["title"] == "Bài 2"
    assert len(subs[1]["units"]) == 2
